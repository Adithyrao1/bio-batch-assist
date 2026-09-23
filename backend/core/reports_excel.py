import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse
from django.utils import timezone
from datetime import datetime
from django.db.models import Sum
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status
from celery.result import AsyncResult
from io import BytesIO

from .models import (
    User, Variety, InitiationLog, MultiplicationLog, 
    RootingLog, HardeningLog, TransplantationLog, Expense,
    Chemical, StockSolutionChemicalUsage, ManpowerExpense
)

def build_production_excel(start_date_str, end_date_str, tech_ids, var_ids):
    if not end_date_str:
        end_date_str = timezone.now().date().isoformat()

    try:
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date() if start_date_str else None
    except ValueError:
        raise ValueError("Invalid date format. Use YYYY-MM-DD.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Progress Report"

    # Styles
    bold_font = Font(bold=True)
    title_font = Font(bold=True, size=12)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    green_fill = PatternFill(start_color="C4D79B", end_color="C4D79B", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                         top=Side(style='thin'), bottom=Side(style='thin'))

    # Helper function to apply styles to a cell
    def style_cell(cell, font=None, alignment=None, fill=None, border=None):
        if font: cell.font = font
        if alignment: cell.alignment = alignment
        if fill: cell.fill = fill
        if border: cell.border = border

    # ---------------------------
    # HEADER ROWS
    # ---------------------------
    # Row 1: Title
    ws.merge_cells('A1:X1')
    cell = ws.cell(row=1, column=1, value="DCM Shriram Ltd., Sugar Unit- Loni")
    style_cell(cell, font=title_font, alignment=center_align, border=thin_border)
    for col in range(2, 25): style_cell(ws.cell(row=1, column=col), border=thin_border)

    # Row 2: Subtitle
    ws.merge_cells('A2:X2')
    cell = ws.cell(row=2, column=1, value="Progress Report of Tissue Culture Seedling Production")
    style_cell(cell, font=title_font, alignment=center_align, border=thin_border)
    for col in range(2, 25): style_cell(ws.cell(row=2, column=col), border=thin_border)

    headers = [
        ("Inoculation Person Name", 1),
        ("Varieties", 1),
        # Initiation
        ("Meristem Tissue Inoculated (No. of Bottles)", 1),
        ("Contaminated Tubes (Nos)", 1),
        ("Contamination %", 1),
        # Multiplication
        ("Ist (No. of Bottles)", 1),
        ("IInd (No. of Bottles)", 1),
        ("IIIrd (No. of Bottles)", 1),
        ("IVth (No. of Bottles)", 1),
        ("Vth (No. of Bottles)", 1),
        ("Contaminated Bottles (Nos)", 1),
        ("Contamination %", 1),
        # Rooting
        ("Basal (No. of Bottles)", 1),
        ("Rooting (No. of Bottles)", 1),
        ("Contaminated Bottles (Nos)", 1),
        ("Contamination %", 1),
        # Hardening
        ("Seedling Transplant in Green House (Nos)", 1),
        ("Seedling Dried In Green House (Nos)", 1),
        ("Mortality % in Green House", 1),
        # Transplantation
        ("Seedling Transplant in Field (Nos)", 1),
        ("Seedling Dried in Field (Nos)", 1),
        ("Mortality % in Field", 1)
    ]

    # Groups
    groups = [
        ("Initiation", 3, 5),
        ("Multiplication", 6, 12),
        ("Rooting", 13, 16),
        ("Hardening", 17, 19),
        ("Transplantation", 20, 22)
    ]

    def build_header_section(start_row, section_title):
        # Section Title Row
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=22)
        cell = ws.cell(row=start_row, column=1, value=section_title)
        style_cell(cell, font=title_font, alignment=center_align, fill=yellow_fill, border=thin_border)
        for col in range(2, 23): style_cell(ws.cell(row=start_row, column=col), border=thin_border)

        # Stage Headers
        ws.merge_cells(start_row=start_row+1, start_column=1, end_row=start_row+2, end_column=1)
        ws.merge_cells(start_row=start_row+1, start_column=2, end_row=start_row+2, end_column=2)
        
        for group_name, start_col, end_col in groups:
            ws.merge_cells(start_row=start_row+1, start_column=start_col, end_row=start_row+1, end_column=end_col)
            cell = ws.cell(row=start_row+1, column=start_col, value=group_name)
            style_cell(cell, font=bold_font, alignment=center_align, fill=green_fill, border=thin_border)
            for c in range(start_col+1, end_col+1): style_cell(ws.cell(row=start_row+1, column=c), border=thin_border)

        # Specific Headers
        for idx, (h_name, width) in enumerate(headers):
            col_idx = idx + 1
            if col_idx <= 2:
                cell = ws.cell(row=start_row+1, column=col_idx, value=h_name)
                style_cell(cell, font=bold_font, alignment=center_align, fill=green_fill, border=thin_border)
                style_cell(ws.cell(row=start_row+2, column=col_idx), border=thin_border)
            else:
                cell = ws.cell(row=start_row+2, column=col_idx, value=h_name)
                style_cell(cell, font=bold_font, alignment=center_align, fill=green_fill, border=thin_border)

        # Set column widths
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 15
        for col_idx in range(3, 23):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 12

    def get_data_for_section(date_filter_kwarg):
        users = User.objects.filter(role__in=['technician', 'admin'])
        if tech_ids:
            users = users.filter(id__in=tech_ids.split(','))
            
        varieties = Variety.objects.all()
        if var_ids:
            varieties = varieties.filter(id__in=var_ids.split(','))
        
        data_rows = []
        totals = [0] * 20
        
        for user in users:
            for variety in varieties:
                # Query aggregates
                # Initiation
                init_logs = InitiationLog.objects.filter(technician=user, variety=variety, **date_filter_kwarg)
                init_aggs = init_logs.aggregate(
                    inoc=Sum('bottles_inoculated'),
                    contam=Sum('contaminated_bottles')
                )
                inoc = init_aggs['inoc'] or 0
                init_contam = init_aggs['contam'] or 0
                
                # Multiplication
                mult_logs = MultiplicationLog.objects.filter(technician=user, variety=variety, **date_filter_kwarg)
                mult_cycles = []
                for c in range(1, 6):
                    c_agg = mult_logs.filter(cycle_number=c).aggregate(prod=Sum('bottles_produced'))['prod'] or 0
                    mult_cycles.append(c_agg)
                mult_contam = mult_logs.aggregate(contam=Sum('contaminated_bottles'))['contam'] or 0
                mult_total_prod = sum(mult_cycles)
                
                # Rooting
                root_logs = RootingLog.objects.filter(technician=user, variety=variety, **date_filter_kwarg)
                root_aggs = root_logs.aggregate(
                    basal=Sum('basal_bottles'),
                    rooting=Sum('rooting_bottles'),
                    contam=Sum('contaminated_bottles')
                )
                root_basal = root_aggs['basal'] or 0
                root_rooting = root_aggs['rooting'] or 0
                root_contam = root_aggs['contam'] or 0
                
                # Hardening
                hard_logs = HardeningLog.objects.filter(technician=user, variety=variety, **date_filter_kwarg)
                hard_aggs = hard_logs.aggregate(
                    trans=Sum('seedlings_transplanted'),
                    died=Sum('seedlings_died')
                )
                hard_trans = hard_aggs['trans'] or 0
                hard_died = hard_aggs['died'] or 0
                
                # Transplantation
                trans_logs = TransplantationLog.objects.filter(technician=user, variety=variety, **date_filter_kwarg)
                trans_aggs = trans_logs.aggregate(
                    trans=Sum('seedlings_transplanted'),
                    died=Sum('seedlings_died')
                )
                trans_trans = trans_aggs['trans'] or 0
                trans_died = trans_aggs['died'] or 0
                
                if inoc == 0 and init_contam == 0 and mult_total_prod == 0 and mult_contam == 0 and \
                   root_basal == 0 and root_rooting == 0 and root_contam == 0 and \
                   hard_trans == 0 and hard_died == 0 and trans_trans == 0 and trans_died == 0:
                   continue
                
                def calc_perc(part, whole):
                    if whole > 0:
                        return f"{(part / whole) * 100:.2f}"
                    return "#DIV/0!" if part == 0 else "100.00"
                
                row_data = [
                    user.get_full_name() or user.username,
                    variety.code,
                    inoc,
                    init_contam,
                    calc_perc(init_contam, inoc),
                    *mult_cycles,
                    mult_contam,
                    calc_perc(mult_contam, mult_total_prod),
                    root_basal,
                    root_rooting,
                    root_contam,
                    calc_perc(root_contam, root_basal + root_rooting),
                    hard_trans,
                    hard_died,
                    calc_perc(hard_died, hard_trans),
                    trans_trans,
                    trans_died,
                    calc_perc(trans_died, trans_trans)
                ]
                data_rows.append(row_data)
                
                totals[0] += inoc
                totals[1] += init_contam
                for i in range(5): totals[3 + i] += mult_cycles[i]
                totals[8] += mult_contam
                totals[10] += root_basal
                totals[11] += root_rooting
                totals[12] += root_contam
                totals[14] += hard_trans
                totals[15] += hard_died
                totals[17] += trans_trans
                totals[18] += trans_died
                
        def calc_perc_num(part, whole):
            if whole > 0:
                return f"{(part / whole) * 100:.2f}"
            return "#DIV/0!"
            
        totals[2] = calc_perc_num(totals[1], totals[0])
        totals[9] = calc_perc_num(totals[8], sum(totals[3:8]))
        totals[13] = calc_perc_num(totals[12], totals[10] + totals[11])
        totals[16] = calc_perc_num(totals[15], totals[14])
        totals[19] = calc_perc_num(totals[18], totals[17])
        
        return data_rows, totals

    current_row = 3
    if start_date and start_date != end_date:
        section_title = f"Selected Period: {start_date_str} to {end_date_str}"
        period_filter = {'date__gte': start_date, 'date__lte': end_date}
    else:
        section_title = f"On Date: {end_date_str}"
        period_filter = {'date': end_date}

    build_header_section(current_row, section_title)
    current_row += 3
    
    on_date_data, on_date_totals = get_data_for_section(period_filter)
    for r_data in on_date_data:
        for c_idx, val in enumerate(r_data):
            cell = ws.cell(row=current_row, column=c_idx+1, value=val)
            style_cell(cell, border=thin_border, alignment=Alignment(horizontal='center'))
            if c_idx < 2:
                style_cell(cell, alignment=Alignment(horizontal='left'))
        current_row += 1
        
    ws.cell(row=current_row, column=1, value="Total")
    ws.cell(row=current_row, column=2, value="")
    style_cell(ws.cell(row=current_row, column=1), font=bold_font, border=thin_border)
    style_cell(ws.cell(row=current_row, column=2), border=thin_border)
    for c_idx, val in enumerate(on_date_totals):
        cell = ws.cell(row=current_row, column=c_idx+3, value=val)
        style_cell(cell, font=bold_font, border=thin_border, alignment=Alignment(horizontal='center'))
    current_row += 2

    build_header_section(current_row, f"To Date (Up to {end_date_str})")
    current_row += 3
    
    to_date_data, to_date_totals = get_data_for_section({'date__lte': end_date})
    for r_data in to_date_data:
        for c_idx, val in enumerate(r_data):
            cell = ws.cell(row=current_row, column=c_idx+1, value=val)
            style_cell(cell, border=thin_border, alignment=Alignment(horizontal='center'))
            if c_idx < 2:
                style_cell(cell, alignment=Alignment(horizontal='left'))
        current_row += 1
        
    ws.cell(row=current_row, column=1, value="Total")
    ws.cell(row=current_row, column=2, value="")
    style_cell(ws.cell(row=current_row, column=1), font=bold_font, border=thin_border)
    style_cell(ws.cell(row=current_row, column=2), border=thin_border)
    for c_idx, val in enumerate(to_date_totals):
        cell = ws.cell(row=current_row, column=c_idx+3, value=val)
        style_cell(cell, font=bold_font, border=thin_border, alignment=Alignment(horizontal='center'))
    
    output = BytesIO()
    wb.save(output)
    file_bytes = output.getvalue()
    filename_date = f"{start_date_str}_to_{end_date_str}" if start_date and start_date != end_date else end_date_str
    filename = f"Progress_Report_{filename_date}.xlsx"
    return file_bytes, filename

def build_expenses_excel(start_date_str, end_date_str):
    expenses = Expense.objects.select_related('category', 'recorded_by').all().order_by('-date')

    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
            expenses = expenses.filter(date__gte=start_date)
        except ValueError:
            raise ValueError("Invalid start_date format. Use YYYY-MM-DD.")
            
    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
            expenses = expenses.filter(date__lte=end_date)
        except ValueError:
            raise ValueError("Invalid end_date format. Use YYYY-MM-DD.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expense Report"

    # Styles
    bold_font = Font(bold=True)
    title_font = Font(bold=True, size=14, color="FFFFFF")
    subtitle_font = Font(bold=True, size=12, color="FFFFFF")
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    header_fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    subheader_fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
    column_header_fill = PatternFill(start_color="FDE68A", end_color="FDE68A", fill_type="solid")
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                         top=Side(style='thin'), bottom=Side(style='thin'))

    def style_cell(cell, font=None, alignment=None, fill=None, border=None):
        if font: cell.font = font
        if alignment: cell.alignment = alignment
        if fill: cell.fill = fill
        if border: cell.border = border

    ws.merge_cells('A1:F1')
    cell = ws.cell(row=1, column=1, value="DCM Shriram Ltd., Sugar Unit- Loni")
    style_cell(cell, font=title_font, alignment=center_align, fill=header_fill, border=thin_border)
    for col in range(2, 7): style_cell(ws.cell(row=1, column=col), fill=header_fill, border=thin_border)

    ws.merge_cells('A2:F2')
    cell = ws.cell(row=2, column=1, value="Expense Report of Tissue Culture Seedling Production")
    style_cell(cell, font=subtitle_font, alignment=center_align, fill=subheader_fill, border=thin_border)
    for col in range(2, 7): style_cell(ws.cell(row=2, column=col), fill=subheader_fill, border=thin_border)

    ws.merge_cells('A3:F3')
    date_text = "All Records"
    if start_date_str and end_date_str:
        date_text = f"Period: {start_date_str} to {end_date_str}"
    elif start_date_str:
        date_text = f"From: {start_date_str}"
    elif end_date_str:
        date_text = f"Up to: {end_date_str}"
        
    cell = ws.cell(row=3, column=1, value=date_text)
    style_cell(cell, font=bold_font, alignment=center_align, fill=PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"), border=thin_border)
    for col in range(2, 7): style_cell(ws.cell(row=3, column=col), border=thin_border)

    headers = ["Date", "Category", "Description", "Invoice / Ref #", "Recorded By", "Amount (₹)"]
    for idx, h in enumerate(headers):
        col_idx = idx + 1
        cell = ws.cell(row=5, column=col_idx, value=h)
        style_cell(cell, font=bold_font, alignment=center_align, fill=column_header_fill, border=thin_border)

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 15

    current_row = 6
    total_amount = 0
    for exp in expenses:
        ws.cell(row=current_row, column=1, value=exp.date.strftime("%Y-%m-%d"))
        ws.cell(row=current_row, column=2, value=exp.category.name if exp.category else "-")
        ws.cell(row=current_row, column=3, value=exp.description)
        ws.cell(row=current_row, column=4, value=exp.invoice_reference)
        ws.cell(row=current_row, column=5, value=exp.recorded_by.get_full_name() or exp.recorded_by.username if exp.recorded_by else "-")
        
        amount_cell = ws.cell(row=current_row, column=6, value=float(exp.amount))
        amount_cell.number_format = '#,##0.00'
        
        for c in range(1, 7):
            align = center_align if c in [1, 4, 5] else (left_align if c in [2, 3] else Alignment(horizontal='right'))
            style_cell(ws.cell(row=current_row, column=c), alignment=align, border=thin_border)
            
        total_amount += float(exp.amount)
        current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    cell = ws.cell(row=current_row, column=1, value="TOTAL EXPENSES")
    style_cell(cell, font=bold_font, alignment=Alignment(horizontal='right'), fill=column_header_fill, border=thin_border)
    for col in range(2, 6): style_cell(ws.cell(row=current_row, column=col), border=thin_border)
    
    total_cell = ws.cell(row=current_row, column=6, value=total_amount)
    total_cell.number_format = '#,##0.00'
    style_cell(total_cell, font=bold_font, alignment=Alignment(horizontal='right'), fill=column_header_fill, border=thin_border)

    ws2 = wb.create_sheet(title="Chemical Expenses")

    chem_header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    chem_subheader_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    chem_column_header_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

    ws2.merge_cells('A1:E1')
    cell = ws2.cell(row=1, column=1, value="DCM Shriram Ltd., Sugar Unit- Loni")
    style_cell(cell, font=title_font, alignment=center_align, fill=chem_header_fill, border=thin_border)
    for col in range(2, 6): style_cell(ws2.cell(row=1, column=col), fill=chem_header_fill, border=thin_border)

    ws2.merge_cells('A2:E2')
    cell = ws2.cell(row=2, column=1, value="Chemical Consumption Cost Report")
    style_cell(cell, font=subtitle_font, alignment=center_align, fill=chem_subheader_fill, border=thin_border)
    for col in range(2, 6): style_cell(ws2.cell(row=2, column=col), fill=chem_subheader_fill, border=thin_border)

    ws2.merge_cells('A3:E3')
    cell = ws2.cell(row=3, column=1, value=date_text)
    style_cell(cell, font=bold_font, alignment=center_align, fill=PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid"), border=thin_border)
    for col in range(2, 6): style_cell(ws2.cell(row=3, column=col), border=thin_border)

    chem_headers = ["Chemical Name", "Unit", "Price per Unit (₹)", "Consumed Quantity", "Total Cost (₹)"]
    for idx, h in enumerate(chem_headers):
        col_idx = idx + 1
        cell = ws2.cell(row=5, column=col_idx, value=h)
        style_cell(cell, font=bold_font, alignment=center_align, fill=chem_column_header_fill, border=thin_border)

    ws2.column_dimensions['A'].width = 35
    ws2.column_dimensions['B'].width = 15
    ws2.column_dimensions['C'].width = 20
    ws2.column_dimensions['D'].width = 20
    ws2.column_dimensions['E'].width = 20

    current_row = 6
    total_chem_cost = 0

    chemicals = Chemical.objects.all().order_by('name')
    for chem in chemicals:
        if start_date_str or end_date_str:
            usages = StockSolutionChemicalUsage.objects.filter(chemical=chem)
            if start_date_str:
                usages = usages.filter(preparation__date__gte=start_date)
            if end_date_str:
                usages = usages.filter(preparation__date__lte=end_date)
            consumed_qty = usages.aggregate(total=Sum('quantity_consumed'))['total'] or 0
        else:
            consumed_qty = float(chem.quantity) - float(chem.remaining_stock)

        if float(consumed_qty) == 0:
            continue

        cost = float(consumed_qty) * float(chem.unit_price)

        ws2.cell(row=current_row, column=1, value=chem.name)
        ws2.cell(row=current_row, column=2, value=chem.unit)
        
        price_cell = ws2.cell(row=current_row, column=3, value=float(chem.unit_price))
        price_cell.number_format = '#,##0.00'
        
        qty_cell = ws2.cell(row=current_row, column=4, value=float(consumed_qty))
        qty_cell.number_format = '#,##0.00'
        
        cost_cell = ws2.cell(row=current_row, column=5, value=cost)
        cost_cell.number_format = '#,##0.00'
        
        for c in range(1, 6):
            align = left_align if c == 1 else (center_align if c == 2 else Alignment(horizontal='right'))
            style_cell(ws2.cell(row=current_row, column=c), alignment=align, border=thin_border)
            
        total_chem_cost += cost
        current_row += 1

    ws2.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    cell = ws2.cell(row=current_row, column=1, value="TOTAL CHEMICAL EXPENSES")
    style_cell(cell, font=bold_font, alignment=Alignment(horizontal='right'), fill=chem_column_header_fill, border=thin_border)
    for col in range(2, 5): style_cell(ws2.cell(row=current_row, column=col), border=thin_border)
    
    total_cell = ws2.cell(row=current_row, column=5, value=total_chem_cost)
    total_cell.number_format = '#,##0.00'
    style_cell(total_cell, font=bold_font, alignment=Alignment(horizontal='right'), fill=chem_column_header_fill, border=thin_border)

    # ---------------------------
    # SHEET 3: MANPOWER EXPENSES
    # ---------------------------
    ws3 = wb.create_sheet(title="Manpower Expenses")
    manpower_header_fill = PatternFill(start_color="8B5CF6", end_color="8B5CF6", fill_type="solid")
    manpower_subheader_fill = PatternFill(start_color="A78BFA", end_color="A78BFA", fill_type="solid")
    manpower_column_header_fill = PatternFill(start_color="DDD6FE", end_color="DDD6FE", fill_type="solid")

    ws3.merge_cells('A1:D1')
    cell = ws3.cell(row=1, column=1, value="DCM Shriram Ltd., Sugar Unit- Loni")
    style_cell(cell, font=title_font, alignment=center_align, fill=manpower_header_fill, border=thin_border)
    for col in range(2, 5): style_cell(ws3.cell(row=1, column=col), fill=manpower_header_fill, border=thin_border)

    ws3.merge_cells('A2:D2')
    cell = ws3.cell(row=2, column=1, value="Manpower Expenses Report")
    style_cell(cell, font=subtitle_font, alignment=center_align, fill=manpower_subheader_fill, border=thin_border)
    for col in range(2, 5): style_cell(ws3.cell(row=2, column=col), fill=manpower_subheader_fill, border=thin_border)

    ws3.merge_cells('A3:D3')
    cell = ws3.cell(row=3, column=1, value="Current Salary Rates")
    style_cell(cell, font=bold_font, alignment=center_align, fill=PatternFill(start_color="F5F3FF", end_color="F5F3FF", fill_type="solid"), border=thin_border)
    for col in range(2, 5): style_cell(ws3.cell(row=3, column=col), border=thin_border)

    manpower_headers = ["Technician", "Notes", "Monthly Salary (₹)", "Daily Rate (₹)"]
    for idx, h in enumerate(manpower_headers):
        col_idx = idx + 1
        cell = ws3.cell(row=5, column=col_idx, value=h)
        style_cell(cell, font=bold_font, alignment=center_align, fill=manpower_column_header_fill, border=thin_border)

    ws3.column_dimensions['A'].width = 30
    ws3.column_dimensions['B'].width = 40
    ws3.column_dimensions['C'].width = 20
    ws3.column_dimensions['D'].width = 20

    current_row = 6
    total_monthly_salary = 0
    total_daily_rate = 0

    manpower_expenses = ManpowerExpense.objects.select_related('technician').all()
    for mp in manpower_expenses:
        ws3.cell(row=current_row, column=1, value=mp.technician.get_full_name() or mp.technician.username)
        ws3.cell(row=current_row, column=2, value=mp.notes)
        
        salary_cell = ws3.cell(row=current_row, column=3, value=float(mp.monthly_salary))
        salary_cell.number_format = '#,##0.00'
        
        rate_cell = ws3.cell(row=current_row, column=4, value=float(mp.daily_rate))
        rate_cell.number_format = '#,##0.00'
        
        for c in range(1, 5):
            align = left_align if c in [1, 2] else Alignment(horizontal='right')
            style_cell(ws3.cell(row=current_row, column=c), alignment=align, border=thin_border)
            
        total_monthly_salary += float(mp.monthly_salary)
        total_daily_rate += float(mp.daily_rate)
        current_row += 1

    ws3.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
    cell = ws3.cell(row=current_row, column=1, value="TOTAL")
    style_cell(cell, font=bold_font, alignment=Alignment(horizontal='right'), fill=manpower_column_header_fill, border=thin_border)
    style_cell(ws3.cell(row=current_row, column=2), border=thin_border)
    
    total_month_cell = ws3.cell(row=current_row, column=3, value=total_monthly_salary)
    total_month_cell.number_format = '#,##0.00'
    style_cell(total_month_cell, font=bold_font, alignment=Alignment(horizontal='right'), fill=manpower_column_header_fill, border=thin_border)

    total_daily_cell = ws3.cell(row=current_row, column=4, value=total_daily_rate)
    total_daily_cell.number_format = '#,##0.00'
    style_cell(total_daily_cell, font=bold_font, alignment=Alignment(horizontal='right'), fill=manpower_column_header_fill, border=thin_border)

    output = BytesIO()
    wb.save(output)
    file_bytes = output.getvalue()
    filename_date = f"{start_date_str}_to_{end_date_str}" if start_date_str and end_date_str else (end_date_str or start_date_str or "All")
    filename = f"Expense_Report_{filename_date}.xlsx"
    return file_bytes, filename


# ── API Views ─────────────────────────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def generate_production_excel_report(request):
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date') or request.GET.get('date')
    tech_ids = request.GET.get('technicians')
    var_ids = request.GET.get('varieties')

    from core.tasks import generate_production_excel_task
    task = generate_production_excel_task.delay(start_date_str, end_date_str, tech_ids, var_ids)
    
    return Response({'task_id': task.id}, status=status.HTTP_202_ACCEPTED)

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def generate_expenses_excel_report(request):
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    from core.tasks import generate_expenses_excel_task
    task = generate_expenses_excel_task.delay(start_date_str, end_date_str)
    
    return Response({'task_id': task.id}, status=status.HTTP_202_ACCEPTED)

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def check_report_status(request, task_id):
    res = AsyncResult(task_id)
    return Response({
        'task_id': task_id,
        'status': res.status,
    })

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def download_report_result(request, task_id):
    from django.core.cache import cache
    import base64
    from django.http import FileResponse, HttpResponse
    
    cache_key = f"report:excel:{task_id}"
    cached_data = cache.get(cache_key)
    
    if not cached_data:
        res = AsyncResult(task_id)
        if res.status == 'SUCCESS':
            return HttpResponse("Report expired. Please generate a new one.", status=410)
        elif res.status == 'FAILURE':
            return HttpResponse(f"Report generation failed: {res.result}", status=500)
        else:
            return HttpResponse("Report is still generating or not found.", status=404)
            
    try:
        file_data = base64.b64decode(cached_data['data'])
        buffer = BytesIO(file_data)
        buffer.seek(0)
        
        response = FileResponse(
            buffer,
            as_attachment=True,
            filename=cached_data['filename']
        )
        response['Content-Type'] = cached_data['content_type']
        
        # Clean up cache immediately
        cache.delete(cache_key)
        return response
    except Exception as e:
        return HttpResponse(f"Failed to process cached report: {str(e)}", status=500)
